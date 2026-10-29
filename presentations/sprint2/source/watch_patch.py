"""Turn a copy of the team's bracelet Google Slides deck into the Afterglow watch deck.
Compares the bracelet and watch GS-edition pptx files shape by shape, matches shapes to the copy's
elements by position, and emits targeted Slides API requests (text edits, image swaps, slide 6 rebuild,
video links). Prints anything the team edited that conflicts, so it can be handled by hand."""
import difflib, hashlib, json, re, sys
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

EMU = 914400
B = Presentation('out/Sprint2_Presentation_Firefly_v2_Bracelet_GoogleSlides.pptx')
W = Presentation('out/Sprint2_Presentation_Afterglow_Watch_GoogleSlides.pptx')
G = json.load(open('watchcopy_read.json'))['content']
GSW = json.load(open('gs_watch.json'))
RAW = 'https://raw.githubusercontent.com/wendyxlin/firefly-band/main/presentations/sprint2/google-slides/assets/'
VURL = 'https://github.com/wendyxlin/firefly-band/blob/main/presentations/sprint2/media/watch-{}.mp4'
VIDEO = {9: 'film-30s', 10: 'app-demo', 11: 'site-demo'}

def u16(s): return len(s.encode('utf-16-le')) // 2
def gtext(e):
    return ''.join(t.get('textRun', {}).get('content', '') for t in e.get('shape', {}).get('text', {}).get('textElements', []))
def gbox(e):
    t = e.get('transform', {}); sz = e.get('size', {})
    return (t.get('translateX', 0), t.get('translateY', 0),
            sz.get('width', {}).get('magnitude', 0) * t.get('scaleX', 1), sz.get('height', {}).get('magnitude', 0) * t.get('scaleY', 1))
def ptext(sh): return '\n'.join(''.join(r.text for r in p.runs) for p in sh.text_frame.paragraphs)
def tokens(s): return re.findall(r'\w+|\s+|[^\w\s]', s)

def match(sh, els, used):
    best, bd = None, 1e18
    for e in els:
        if e['objectId'] in used: continue
        x, y, w, h = gbox(e)
        d = abs(x - sh.left) + abs(y - sh.top) + abs(w - sh.width) + abs(h - sh.height)
        if d < bd: best, bd = e, d
    return best if bd < 0.08 * EMU else None

def text_edits(oid, old, new):
    """Requests turning `old` into `new` inside shape oid, word-level, applied back to front.
    Replacements insert after the first replaced character so the new text keeps that character's style."""
    a, b = tokens(old), tokens(new)
    offs = [0]
    for t in a: offs.append(offs[-1] + u16(t))
    ops = [list(o) for o in difflib.SequenceMatcher(None, a, b, autojunk=False).get_opcodes()]
    # merge changes separated by short unchanged stretches into one phrase-level replace
    merged = []
    for o in ops:
        if merged and merged[-1][0] != 'equal' and o[0] != 'equal':
            merged[-1][2], merged[-1][4] = o[2], o[4]; merged[-1][0] = 'replace'
        elif o[0] == 'equal' and merged and merged[-1][0] != 'equal' and len(''.join(a[o[1]:o[2]])) < 14 and o is not ops[-1] and '\n' not in ''.join(a[o[1]:o[2]]):
            merged[-1][2], merged[-1][4] = o[2], o[4]; merged[-1][0] = 'replace'
        else:
            merged.append(o)
    R = []
    for op, i1, i2, j1, j2 in reversed(merged):
        if op == 'equal': continue
        s, e = offs[i1], offs[i2]; ins = ''.join(b[j1:j2])
        if op == 'replace' and not ins: op = 'delete'
        if op == 'replace' and s == e: op = 'insert'
        if op == 'delete':
            R.append({'deleteText': {'objectId': oid, 'textRange': {'type': 'FIXED_RANGE', 'startIndex': s, 'endIndex': e}}})
        elif op == 'insert':
            R.append({'insertText': {'objectId': oid, 'insertionIndex': s, 'text': ins}})
        else:
            n = u16(ins)
            R.append({'insertText': {'objectId': oid, 'insertionIndex': s + 1, 'text': ins}})
            if e + n > s + 1 + n:
                R.append({'deleteText': {'objectId': oid, 'textRange': {'type': 'FIXED_RANGE', 'startIndex': s + 1 + n, 'endIndex': e + n}}})
            R.append({'deleteText': {'objectId': oid, 'textRange': {'type': 'FIXED_RANGE', 'startIndex': s, 'endIndex': s + 1}}})
    return R

reqs, report = [], []
for si, (bs, ws, gs) in enumerate(zip(B.slides, W.slides, G['slides']), 1):
    pid = gs['objectId']; els = gs.get('pageElements', []); used = set()
    bsh, wsh = list(bs.shapes), list(ws.shapes)
    if si == 6:  # different layout: 3 cards -> 2 cards. Keep the title, rebuild the rest from the watch deck.
        title = match(bsh[0], els, used); used.add(title['objectId'])
        for e in els:
            if e['objectId'] != title['objectId']: reqs.append({'deleteObject': {'objectId': e['objectId']}})
        cur, want = gtext(title).rstrip('\n'), ptext(wsh[0])
        if cur != want: reqs += text_edits(title['objectId'], cur, want)
        src = GSW['slides'][5]; old = 'Ws06'
        for r in src:
            k, v = next(iter(r.items()))
            oid = v.get('objectId', '')
            if k in ('createSlide', 'updatePageProperties') or any(oid.endswith(x) for x in ('_hdr', '_dots')) or '_ln' in oid or oid.endswith('_e00'):
                continue
            r = json.loads(json.dumps(r).replace(json.dumps(src[0]['createSlide']['objectId'])[1:-1], 'aw06'))
            k, v = next(iter(r.items()))
            if 'elementProperties' in v: v['elementProperties']['pageObjectId'] = pid
            reqs.append(r)
        report.append(f'slide 6 rebuilt ({len(els)-1} deleted)')
        continue
    for k, (b, w) in enumerate(zip(bsh, wsh)):
        e = match(b, els, used)
        if b.shape_type == MSO_SHAPE_TYPE.PICTURE:
            same = hashlib.sha1(b.image.blob).digest() == hashlib.sha1(w.image.blob).digest()
            if e is None:
                if not same: report.append(f'slide {si}: picture {k} changed but no element found')
                continue
            used.add(e['objectId'])
            if not same:
                fn = hashlib.sha1(w.image.blob).hexdigest()[:16] + '.' + w.image.ext.replace('jpeg', 'jpg')
                reqs.append({'replaceImage': {'imageObjectId': e['objectId'], 'url': RAW + fn, 'imageReplaceMethod': 'CENTER_CROP'}})
            nxt = bsh[k + 1] if k + 1 < len(bsh) else None
            if si in VIDEO and nxt is not None and nxt.has_text_frame and nxt.text_frame.text.strip() == '▶':
                reqs.append({'updateImageProperties': {'objectId': e['objectId'], 'imageProperties': {'link': {'url': VURL.format(VIDEO[si])}}, 'fields': 'link'}})
            continue
        if not (b.has_text_frame and w.has_text_frame): continue
        bt, wt = ptext(b), ptext(w)
        if e is None:
            if bt != wt or bt.strip() == '▶': report.append(f'slide {si}: text shape {k} {bt[:30]!r} not in copy (team removed it?)')
            continue
        used.add(e['objectId'])
        if bt == wt: continue
        cur = gtext(e).rstrip('\n')
        if cur == bt:
            reqs += text_edits(e['objectId'], cur, wt)
        else:
            report.append(f'slide {si}: CONFLICT {e["objectId"]}\n   team: {cur!r}\n   brac: {bt!r}\n   watch:{wt!r}')
    for e in els:
        if 'video' in e:
            reqs.append({'deleteObject': {'objectId': e['objectId']}}); report.append(f'slide {si}: removed bracelet video {e["video"]["id"]}')

json.dump(reqs, open('watch_patch.json', 'w'))
print(len(reqs), 'requests,', len(json.dumps(reqs)), 'bytes')
print('\n'.join(report))

# Title slide: the team cleared the subtitle and moved the tagline; keep their layout, swap the tagline.
t34 = next(e for e in G['slides'][0]['pageElements'] if e['objectId'] == 'h36b404f39f93f03e_3_34')
reqs += text_edits('h36b404f39f93f03e_3_34', gtext(t34).rstrip('\n'), ptext(list(W.slides[0].shapes)[2]))
# Speaker notes: swap where the watch script differs and the team hasn't changed the bracelet one.
# (The copy's notes were read and still match the bracelet scripts word for word.)
for si, (bs, ws, gsl) in enumerate(zip(B.slides, W.slides, G['slides']), 1):
    bn, wn = bs.notes_slide.notes_text_frame.text, ws.notes_slide.notes_text_frame.text
    if bn == wn: continue
    oid = gsl['slideProperties']['notesPage']['notesProperties']['speakerNotesObjectId']
    reqs += [{'deleteText': {'objectId': oid, 'textRange': {'type': 'ALL'}}}, {'insertText': {'objectId': oid, 'insertionIndex': 0, 'text': wn}}]
# Split into batches of about 14 KB, never splitting one shape's text edits.
batches, cur, size, last = [], [], 0, None
for r in reqs:
    v = next(iter(r.values())); oid = v.get('objectId') or v.get('imageObjectId')
    n = len(json.dumps(r))
    if cur and size + n > 14000 and oid != last: batches.append(cur); cur, size = [], 0
    cur.append(r); size += n; last = oid
batches.append(cur)
for i, b in enumerate(batches): json.dump(b, open(f'wb{i}.json', 'w'), ensure_ascii=False)
print(len(reqs), 'total;', [len(b) for b in batches], 'batches')
